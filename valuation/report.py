"""
Builds a formatted Excel workbook from a valuation run: a Summary sheet,
an Assumptions sheet, the year-by-year FCF Projection, and a
WACC x terminal-growth Sensitivity table with a colour scale — the same
shape of workbook an analyst would hand-build in Excel, just generated in
seconds from live data.
"""

from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill, Alignment
from openpyxl.formatting.rule import ColorScaleRule
from openpyxl.utils import get_column_letter
from openpyxl.chart import BarChart, Reference

HEADER_FILL = PatternFill(start_color="1F4E78", end_color="1F4E78", fill_type="solid")
HEADER_FONT = Font(color="FFFFFF", bold=True)
TITLE_FONT = Font(bold=True, size=14)
LABEL_FONT = Font(bold=True)
CURRENCY_FMT = '#,##0.00'
BIG_NUMBER_FMT = '#,##0,,"M"'   # displays raw dollars as millions
PCT_FMT = '0.0%'


def _style_header_row(ws, row, num_cols):
    for col in range(1, num_cols + 1):
        cell = ws.cell(row=row, column=col)
        cell.fill = HEADER_FILL
        cell.font = HEADER_FONT
        cell.alignment = Alignment(horizontal="center")


def _autosize(ws, widths):
    for i, width in enumerate(widths, start=1):
        ws.column_dimensions[get_column_letter(i)].width = width


def _write_summary_sheet(wb, financials, assumptions, result):
    ws = wb.active
    ws.title = "Summary"

    ws["A1"] = f"{financials.name} ({financials.ticker}) — DCF Valuation"
    ws["A1"].font = TITLE_FONT
    ws.merge_cells("A1:B1")

    rows = [
        ("Current share price", financials.current_price, CURRENCY_FMT),
        ("Implied share price (DCF)", result.implied_share_price, CURRENCY_FMT),
        ("Upside / (downside)", result.upside_pct, PCT_FMT),
        ("", None, None),
        ("WACC used", result.wacc, PCT_FMT),
        ("Terminal growth rate", assumptions.terminal_growth_rate, PCT_FMT),
        ("Forecast horizon (years)", assumptions.forecast_years, "0"),
        ("", None, None),
        ("Enterprise value", result.enterprise_value, BIG_NUMBER_FMT),
        ("Net debt", financials.net_debt, BIG_NUMBER_FMT),
        ("Equity value", result.equity_value, BIG_NUMBER_FMT),
        ("Shares outstanding", financials.shares_outstanding, "#,##0,,\"M\""),
        ("", None, None),
        ("PV of explicit-period FCF", sum(result.present_values), BIG_NUMBER_FMT),
        ("PV of terminal value", result.pv_terminal_value, BIG_NUMBER_FMT),
    ]

    r = 3
    for label, value, fmt in rows:
        ws.cell(row=r, column=1, value=label).font = LABEL_FONT
        if value is not None:
            cell = ws.cell(row=r, column=2, value=value)
            if fmt:
                cell.number_format = fmt
        r += 1

    _autosize(ws, [30, 18])


def _write_assumptions_sheet(wb, financials, assumptions, result):
    ws = wb.create_sheet("Assumptions")
    ws["A1"] = "Assumptions"
    ws["A1"].font = TITLE_FONT

    rows = [
        ("Revenue growth rate", assumptions.revenue_growth_rate, PCT_FMT),
        ("EBIT margin", assumptions.ebit_margin, PCT_FMT),
        ("Tax rate", assumptions.tax_rate, PCT_FMT),
        ("D&A (% of revenue)", assumptions.da_pct_revenue, PCT_FMT),
        ("Capex (% of revenue)", assumptions.capex_pct_revenue, PCT_FMT),
        ("NWC change (% of revenue change)", assumptions.nwc_pct_of_revenue_change, PCT_FMT),
        ("", None, None),
        ("Beta", financials.beta, "0.00"),
        ("Risk-free rate", assumptions.risk_free_rate, PCT_FMT),
        ("Market risk premium", assumptions.market_risk_premium, PCT_FMT),
        ("Pre-tax cost of debt", assumptions.cost_of_debt, PCT_FMT),
        ("WACC (computed or overridden)", result.wacc, PCT_FMT),
    ]

    r = 3
    for label, value, fmt in rows:
        ws.cell(row=r, column=1, value=label).font = LABEL_FONT
        if value is not None:
            cell = ws.cell(row=r, column=2, value=value)
            if fmt:
                cell.number_format = fmt
        r += 1

    _autosize(ws, [32, 14])


def _write_projection_sheet(wb, result):
    ws = wb.create_sheet("FCF Projection")
    headers = ["Year", "Revenue", "EBIT", "NOPAT", "D&A", "Capex",
               "Change in NWC", "Unlevered FCF", "PV of FCF"]
    ws.append(headers)
    _style_header_row(ws, 1, len(headers))

    for proj, pv in zip(result.projections, result.present_values):
        ws.append([
            proj.year, proj.revenue, proj.ebit, proj.nopat, proj.da,
            proj.capex, proj.change_in_nwc, proj.unlevered_fcf, pv,
        ])

    for row in ws.iter_rows(min_row=2, min_col=2, max_col=len(headers)):
        for cell in row:
            cell.number_format = BIG_NUMBER_FMT

    _autosize(ws, [8, 16, 14, 14, 12, 12, 14, 16, 14])


def _write_sensitivity_sheet(wb, wacc_values, growth_values, price_grid):
    ws = wb.create_sheet("Sensitivity")
    ws["A1"] = "Implied share price: WACC (rows) vs. terminal growth rate (columns)"
    ws["A1"].font = TITLE_FONT
    ws.merge_cells(start_row=1, start_column=1, end_row=1, end_column=len(growth_values) + 1)

    header_row = 3
    ws.cell(row=header_row, column=1, value="WACC \\ g")
    for j, growth in enumerate(growth_values):
        ws.cell(row=header_row, column=2 + j, value=growth).number_format = PCT_FMT
    _style_header_row(ws, header_row, len(growth_values) + 1)

    for i, wacc in enumerate(wacc_values):
        row_num = header_row + 1 + i
        cell = ws.cell(row=row_num, column=1, value=wacc)
        cell.number_format = PCT_FMT
        cell.font = LABEL_FONT
        for j, price in enumerate(price_grid[i]):
            ws.cell(row=row_num, column=2 + j, value=price)
            if price is not None:
                ws.cell(row=row_num, column=2 + j).number_format = CURRENCY_FMT

    first_data_row = header_row + 1
    last_data_row = header_row + len(wacc_values)
    last_col = 1 + len(growth_values)
    data_range = (
        f"{get_column_letter(2)}{first_data_row}:"
        f"{get_column_letter(last_col)}{last_data_row}"
    )
    ws.conditional_formatting.add(
        data_range,
        ColorScaleRule(
            start_type="min", start_color="F8696B",
            mid_type="percentile", mid_value=50, mid_color="FFEB84",
            end_type="max", end_color="63BE7B",
        ),
    )

    _autosize(ws, [12] + [12] * len(growth_values))


def _write_comps_sheet(wb, financials, result, comps_result):
    """Peer trading multiples, the implied price under each method, and a
    football-field chart comparing them against the DCF and the current
    market price — the standard way analysts show two valuation methods
    side by side rather than presenting a single number as gospel."""
    ws = wb.create_sheet("Comps")
    ws["A1"] = f"{financials.name} ({financials.ticker}) — Comparable Companies"
    ws["A1"].font = TITLE_FONT
    ws.merge_cells("A1:F1")

    headers = ["Ticker", "Name", "P/E", "EV/EBITDA", "EV/Revenue"]
    header_row = 3
    for j, h in enumerate(headers, start=1):
        ws.cell(row=header_row, column=j, value=h)
    _style_header_row(ws, header_row, len(headers))

    r = header_row + 1
    for peer in comps_result.peers:
        ws.cell(row=r, column=1, value=peer.ticker)
        ws.cell(row=r, column=2, value=peer.name)
        ws.cell(row=r, column=3, value=peer.pe).number_format = "0.0x"
        ws.cell(row=r, column=4, value=peer.ev_ebitda).number_format = "0.0x"
        ws.cell(row=r, column=5, value=peer.ev_revenue).number_format = "0.0x"
        r += 1

    summary_rows = [
        ("Average", comps_result.avg_pe, comps_result.avg_ev_ebitda, comps_result.avg_ev_revenue),
        ("Median", comps_result.median_pe, comps_result.median_ev_ebitda, comps_result.median_ev_revenue),
    ]
    for label, pe, ev_ebitda, ev_rev in summary_rows:
        ws.cell(row=r, column=2, value=label).font = LABEL_FONT
        ws.cell(row=r, column=3, value=pe).number_format = "0.0x"
        ws.cell(row=r, column=4, value=ev_ebitda).number_format = "0.0x"
        ws.cell(row=r, column=5, value=ev_rev).number_format = "0.0x"
        r += 1

    r += 1
    ws.cell(row=r, column=1, value="Football field ($/share)").font = LABEL_FONT
    r += 1
    chart_header_row = r
    ws.cell(row=r, column=1, value="Method")
    ws.cell(row=r, column=2, value="Implied price")
    _style_header_row(ws, r, 2)

    methods = [
        ("Current market price", financials.current_price),
        ("DCF", result.implied_share_price),
        ("Comps: EV/EBITDA", comps_result.implied_price_ev_ebitda),
        ("Comps: EV/Revenue", comps_result.implied_price_ev_revenue),
        ("Comps: P/E", comps_result.implied_price_pe),
    ]
    first_chart_data_row = r + 1
    for label, price in methods:
        r += 1
        ws.cell(row=r, column=1, value=label)
        if price is not None:
            ws.cell(row=r, column=2, value=price).number_format = CURRENCY_FMT
    last_chart_data_row = r

    chart = BarChart()
    chart.type = "bar"  # horizontal bars, easiest to read as a football field
    chart.title = "Implied share price by method"
    chart.y_axis.title = None
    chart.x_axis.title = "$ / share"
    data = Reference(ws, min_col=2, min_row=chart_header_row,
                      max_row=last_chart_data_row)
    categories = Reference(ws, min_col=1, min_row=first_chart_data_row,
                            max_row=last_chart_data_row)
    chart.add_data(data, titles_from_data=True)
    chart.set_categories(categories)
    chart.legend = None
    ws.add_chart(chart, f"D{chart_header_row}")

    _autosize(ws, [24, 28, 10, 12, 12])


def build_excel_report(financials, assumptions, result,
                        wacc_values, growth_values, price_grid, output_path,
                        comps_result=None):
    wb = Workbook()
    _write_summary_sheet(wb, financials, assumptions, result)
    _write_assumptions_sheet(wb, financials, assumptions, result)
    _write_projection_sheet(wb, result)
    _write_sensitivity_sheet(wb, wacc_values, growth_values, price_grid)
    if comps_result is not None:
        _write_comps_sheet(wb, financials, result, comps_result)
    wb.save(output_path)
    return output_path
